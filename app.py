import io
import datetime
import streamlit as st
import openpyxl

import engine

st.set_page_config(page_title="FR TRIM Auto-Fill", page_icon="🧵", layout="centered")

# --------------------------------------------------------------------------
# i18n
# --------------------------------------------------------------------------

LANGS = {"ko": "한국어", "en": "English", "vi": "Tiếng Việt"}

TEXT = {
    "ko": {
        "title": "🧵 FR Chart TRIM 자동 채우기",
        "caption": (
            "여러 개의 Balance Sheet(스타일별) + FR(Friday Review) 차트 1개를 올리면, "
            "각 스타일의 TRIM 영역을 Balance Sheet 데이터로 자동으로 채운 FR 파일을 만들어줍니다."
        ),
        "rules_expander": "ℹ️ 처리 규칙 (한번 확인해보세요)",
        "rules_md": """
- Balance Sheet의 **GENERAL TRIM / SPECIAL TRIM** 항목표에서 **Item/Color, Qty Needed, Stock, Shipped Qty, ETD, ETA, AWB#(Tracking#)** 총 7개 항목을 **뜻으로** 찾아서 가져옵니다. 컬럼 순서나 정확한 표기(예: "Tracking#" vs "AWB#", "Shipped Qty" vs "QTY(PCS)", "Qty Needed" vs "QTY NEEDED")가 Balance Sheet마다 달라도 자동으로 인식됩니다.
- 거래처(Mill/Supplier) 이름은 가져오지 않습니다.
- 이름에 **"Hangtag"** 가 들어간 항목은 **사이즈별로 줄을 나눠서** 넣습니다.
- **1st order**든 **2nd order**든, 그 shipment의 ETD/ETA/Tracking#/Shipped Qty **4개 중 하나라도 값이 있을 때만** 해당 라인을 추가합니다 (4개가 전부 비어있으면 그 shipment는 건너뜁니다). Qty Needed/Stock는 품목 단위 값이라 추가된 라인에 그대로 같이 들어갑니다.
- FR 차트 쪽에서는 헤더에 **"trim"**이라는 단어가 들어간 칸(TRIM, TRIM SCHEDULE 등 무엇이든)을 자동으로 찾아 연동합니다. 이미 만들어져 있는 하위 컬럼(예: AWB#, QTY(PCS))의 데이터는 그대로 보존되고, 7개 항목 중 아직 없는 것(예: Qty Needed, Stock)은 새로 추가되며, **항목 순서는 기존 순서와 관계없이 Item → Qty Needed → Stock → Shipped Qty → ETD → ETA → AWB#(Tracking#) 순서로 항상 맞춰서 재배열**됩니다.
- 모든 라벨/데이터 글자 크기는 20pt로 통일되고, "TRIM SCHEDULE" 같은 상단 배너 칸은 늘어난 7개 컬럼 전체를 덮도록 자동으로 넓어집니다. 다른 컬럼(예: FB TEST, TRIM TEST 등)의 헤더 영역도 끊김 없이 같은 줄 높이로 이어지도록 자동으로 병합됩니다.
- 스타일당 TRIM 영역을 **7개 컬럼(Item / Qty Needed / Stock / Shipped Qty / ETD / ETA / AWB#(Tracking#)) x N개 행**의 실제 표(테두리 포함)로 바꿔 넣습니다.
- 숫자 항목(Qty Needed/Stock/Shipped Qty)에 `#REF!` 같은 엑셀 오류 값이 있으면 제외하고 숫자만 가져오며, ETD/ETA는 실제 날짜 값만 가져옵니다.
- 처음 처리되는 스타일의 경우, 기존에 손으로 적혀있던 내용은 새 데이터 아래로 내려가 보존되며(삭제되지 않음) — 내용이 한 가지 항목뿐인 경우 보기 편하게 7개 칸을 병합해서 넣습니다. 이렇게 보존된 내용은 이후 다시 처리해도 계속 유지됩니다.
- 필요한 줄 수가 기존 스타일 블록의 행 수보다 많으면, 해당 스타일 블록 전체를 자동으로 늘리고 (다른 컬럼들의 병합, 수식, 아래 스타일들 위치도 같이 안전하게 조정) 부족한 만큼 채웁니다.
- 헤더에 **BULK FB / TRIM** 그룹 제목과 구분선, 그리고 7개 컬럼 라벨을 자동으로 넣습니다. 라벨을 넣을 별도의 줄이 없는 파일이라면, 줄 하나를 새로 만들어서 라벨 전용으로 씁니다 (기존 데이터는 한 줄 아래로 안전하게 이동).
""",
        "split_keyword_label": "사이즈별로 줄을 나눌 항목 키워드 (기본: hangtag)",
        "step1_header": "1. Balance Sheet 업로드 (여러 개 선택 가능, 스타일당 1개)",
        "bs_uploader_label": "Balance Sheet 파일들",
        "step2_header": "2. FR (Friday Review) 차트 업로드 (1개)",
        "fr_uploader_label": "FR 차트 파일",
        "sheet_select_label": "FR 파일 안에서 사용할 시트를 선택하세요",
        "fr_read_error": "FR 파일을 읽는 중 문제가 발생했습니다: {e}",
        "process_button": "🚀 처리 시작",
        "spinner": "처리 중입니다...",
        "process_error": "처리 중 오류가 발생했습니다: {e}",
        "results_header": "처리 결과",
        "result_ok": "스타일 {style}: {lines}줄 반영 완료 (행 {start}~{end}{added}{preserved})",
        "rows_added_suffix": ", {n}행 추가",
        "rows_preserved_suffix": ", 기존에 적혀있던 {n}줄은 지우지 않고 아래로 내림",
        "result_warn": "스타일 {style} ({sheet}): {status}",
        "status_no_style": "Balance Sheet에서 스타일 번호를 찾지 못했습니다",
        "status_no_trim_section": "GENERAL TRIM 섹션을 찾지 못했습니다",
        "status_style_not_found_in_fr": "FR 차트에서 이 스타일을 찾지 못했습니다",
        "status_fr_columns_not_found": "FR 차트에서 STYLE NO / TRIM 컬럼을 찾지 못했습니다",
        "download_button": "⬇️ 업데이트된 FR 파일 다운로드",
        "lang_label": "Language / 언어 / Ngôn ngữ",
    },
    "en": {
        "title": "🧵 FR Chart TRIM Auto-Fill",
        "caption": (
            "Upload multiple Balance Sheets (one per style) plus one FR (Friday Review) chart, "
            "and this will automatically fill in each style's TRIM section using the Balance Sheet data."
        ),
        "rules_expander": "ℹ️ How it works (worth a quick read)",
        "rules_md": """
- Pulls 7 fields by **meaning**, not fixed position, from the Balance Sheet's **GENERAL TRIM / SPECIAL TRIM** item table: **Item/Color, Qty Needed, Stock, Shipped Qty, ETD, ETA, AWB#(Tracking#)**. Column order and exact wording can differ between files ("Tracking#" vs "AWB#", "Shipped Qty" vs "QTY (PCS)", "Qty Needed" vs "QTY NEEDED") and it still matches correctly.
- Vendor / Mill-Supplier name is **not** included, to keep it as condensed as possible.
- Any item whose name contains **"Hangtag"** is **split into one line per size**.
- For **either** the 1st or 2nd order leg, a line is only added when **at least one** of that leg's ETD/ETA/Tracking#/Shipped Qty is filled in - skipped entirely if all four are blank. Qty Needed/Stock are per-item, so they ride along on whichever lines do get added.
- On the FR chart side, any header cell whose text contains **"trim"** (TRIM, TRIM SCHEDULE, or anything else with that word) is found automatically. Data already in existing sub-columns (e.g. AWB#, QTY (PCS)) is kept, whichever of the 7 fields are still missing (e.g. Qty Needed, Stock) get added, and **the columns are always rearranged into the canonical order** - Item, Qty Needed, Stock, Shipped Qty, ETD, ETA, AWB#(Tracking#) - regardless of what order they were in before.
- All label/data text is set to a uniform 20pt, and the banner cell above the grid (e.g. "TRIM SCHEDULE") automatically widens to span all 7 columns. Other nearby columns' header areas (e.g. FB TEST, TRIM TEST) are extended to match the same row height so nothing looks disconnected.
- The FR chart's TRIM area for each style is unmerged and rebuilt as a real bordered table: **7 columns (Item / Qty Needed / Stock / Shipped Qty / ETD / ETA / AWB#(Tracking#)) x N rows**.
- Excel error values (like `#REF!`) in quantity fields are dropped rather than pulled in - only real numbers are kept. ETD/ETA only pull genuine date values.
- The first time a style is ever processed, any hand-typed content already there is moved below the new rows instead of being deleted - a single freeform note gets merged across all 7 columns for readability. That preserved content keeps being carried forward on every later re-run too, so it's never lost.
- If more rows are needed than the style's current row block has, the whole block is safely grown (other merged columns, formulas, and the position of styles below are all adjusted automatically).
- Header banners **BULK FB / TRIM** with a divider, plus all 7 column labels, are set up automatically. If a chart has no row of its own for these labels, one is inserted so the labels get a permanent home (existing data shifts down by one row safely).
""",
        "split_keyword_label": "Keyword for items to split by size (default: hangtag)",
        "step1_header": "1. Upload Balance Sheets (multiple allowed, one per style)",
        "bs_uploader_label": "Balance Sheet files",
        "step2_header": "2. Upload the FR (Friday Review) chart (one file)",
        "fr_uploader_label": "FR chart file",
        "sheet_select_label": "Choose which sheet to use in the FR file",
        "fr_read_error": "There was a problem reading the FR file: {e}",
        "process_button": "🚀 Start processing",
        "spinner": "Processing...",
        "process_error": "An error occurred while processing: {e}",
        "results_header": "Results",
        "result_ok": "Style {style}: {lines} line(s) written (rows {start}-{end}{added}{preserved})",
        "rows_added_suffix": ", {n} row(s) added",
        "rows_preserved_suffix": ", kept {n} pre-existing line(s) below (not deleted)",
        "result_warn": "Style {style} ({sheet}): {status}",
        "status_no_style": "Could not find a style number in this Balance Sheet",
        "status_no_trim_section": "Could not find a GENERAL TRIM section",
        "status_style_not_found_in_fr": "This style was not found in the FR chart",
        "status_fr_columns_not_found": "Could not find the STYLE NO / TRIM columns in the FR chart",
        "download_button": "⬇️ Download updated FR file",
        "lang_label": "Language / 언어 / Ngôn ngữ",
    },
    "vi": {
        "title": "🧵 Tự động điền TRIM cho FR Chart",
        "caption": (
            "Tải lên nhiều Balance Sheet (mỗi style một file) cùng với 1 file FR (Friday Review) chart, "
            "công cụ sẽ tự động điền phần TRIM của từng style bằng dữ liệu từ Balance Sheet."
        ),
        "rules_expander": "ℹ️ Quy tắc xử lý (nên xem qua một lần)",
        "rules_md": """
- Lấy 7 thông tin theo **ý nghĩa**, không theo vị trí cố định, từ bảng mục **GENERAL TRIM / SPECIAL TRIM** trong Balance Sheet: **Item/Color, Qty Needed, Stock, Shipped Qty, ETD, ETA, AWB#(Tracking#)**. Thứ tự cột và cách ghi chính xác có thể khác nhau giữa các file ("Tracking#" vs "AWB#", "Shipped Qty" vs "QTY (PCS)", "Qty Needed" vs "QTY NEEDED") nhưng vẫn tự nhận diện đúng.
- **Không** lấy tên nhà cung cấp (Mill/Supplier) để giữ nội dung gọn nhất có thể.
- Các mục có tên chứa **"Hangtag"** sẽ được **tách riêng theo từng size**.
- Với **1st order** hoặc **2nd order**, dòng chỉ được thêm vào khi **có ít nhất một** trong 4 giá trị ETD/ETA/Tracking#/Shipped Qty của shipment đó được điền (nếu cả 4 đều trống thì bỏ qua shipment đó). Qty Needed/Stock là giá trị theo item nên sẽ đi kèm theo các dòng được thêm vào.
- Ở phía FR chart, bất kỳ ô tiêu đề nào có chữ **"trim"** (TRIM, TRIM SCHEDULE, hoặc bất kỳ chữ nào có từ đó) đều được tự động tìm ra. Dữ liệu đã có trong các cột con (ví dụ AWB#, QTY (PCS)) vẫn được giữ lại, những cột nào trong 7 cột còn thiếu (ví dụ Qty Needed, Stock) sẽ được thêm mới, và **thứ tự cột luôn được sắp xếp lại** theo đúng thứ tự chuẩn: Item, Qty Needed, Stock, Shipped Qty, ETD, ETA, AWB#(Tracking#), bất kể thứ tự ban đầu.
- Toàn bộ chữ (nhãn và dữ liệu) được chỉnh về cỡ 20pt thống nhất, và ô banner phía trên (ví dụ "TRIM SCHEDULE") sẽ tự động mở rộng để che hết 7 cột. Các cột khác gần đó (ví dụ FB TEST, TRIM TEST) cũng được nối liền theo cùng chiều cao dòng để không bị đứt đoạn.
- Ô TRIM trong FR chart (1 ô gộp cho mỗi style) sẽ được tách ra và dựng lại thành bảng thật có viền: **7 cột (Item / Qty Needed / Stock / Shipped Qty / ETD / ETA / AWB#(Tracking#)) x N dòng**.
- Các giá trị lỗi Excel (như `#REF!`) trong các cột số lượng sẽ bị loại bỏ, chỉ lấy số thật. ETD/ETA chỉ lấy giá trị ngày thật.
- Lần đầu xử lý một style, nội dung viết tay đã có sẽ được đưa xuống dưới các dòng mới (không bị xóa) — nếu chỉ có 1 nội dung duy nhất, nó sẽ được gộp vào 7 cột cho dễ xem. Nội dung này sẽ tiếp tục được giữ lại ở những lần xử lý sau.
- Nếu số dòng cần nhiều hơn số hàng hiện có của khối style đó, toàn bộ khối sẽ được **tự động mở rộng** (các cột gộp khác, công thức, và vị trí của các style phía dưới đều được điều chỉnh an toàn).
- Tiêu đề nhóm **BULK FB / TRIM** cùng đường phân cách, và đủ 7 nhãn cột, sẽ được tự động thêm vào. Nếu chart chưa có dòng riêng cho các nhãn này, một dòng mới sẽ được chèn vào để nhãn có vị trí cố định (dữ liệu cũ sẽ dịch xuống 1 dòng một cách an toàn).
""",
        "split_keyword_label": "Từ khóa để tách dòng theo size (mặc định: hangtag)",
        "step1_header": "1. Tải lên Balance Sheet (có thể chọn nhiều file, mỗi style 1 file)",
        "bs_uploader_label": "Các file Balance Sheet",
        "step2_header": "2. Tải lên FR (Friday Review) chart (1 file)",
        "fr_uploader_label": "File FR chart",
        "sheet_select_label": "Chọn sheet muốn dùng trong file FR",
        "fr_read_error": "Có lỗi khi đọc file FR: {e}",
        "process_button": "🚀 Bắt đầu xử lý",
        "spinner": "Đang xử lý...",
        "process_error": "Đã xảy ra lỗi trong quá trình xử lý: {e}",
        "results_header": "Kết quả xử lý",
        "result_ok": "Style {style}: đã ghi {lines} dòng (hàng {start}-{end}{added}{preserved})",
        "rows_added_suffix": ", đã thêm {n} hàng",
        "rows_preserved_suffix": ", giữ lại {n} dòng cũ phía dưới (không xóa)",
        "result_warn": "Style {style} ({sheet}): {status}",
        "status_no_style": "Không tìm thấy mã style trong Balance Sheet này",
        "status_no_trim_section": "Không tìm thấy phần GENERAL TRIM",
        "status_style_not_found_in_fr": "Không tìm thấy style này trong FR chart",
        "status_fr_columns_not_found": "Không tìm thấy cột STYLE NO / TRIM trong FR chart",
        "download_button": "⬇️ Tải file FR đã cập nhật",
        "lang_label": "Language / 언어 / Ngôn ngữ",
    },
}

if "lang" not in st.session_state:
    st.session_state.lang = "ko"

lang_code = st.selectbox(
    TEXT[st.session_state.lang]["lang_label"],
    options=list(LANGS.keys()),
    format_func=lambda k: LANGS[k],
    index=list(LANGS.keys()).index(st.session_state.lang),
    key="lang_selector",
)
st.session_state.lang = lang_code
T = TEXT[lang_code]


def status_text(code):
    return T.get(f"status_{code}", code)


# --------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------

st.title(T["title"])
st.caption(T["caption"])

with st.expander(T["rules_expander"], expanded=False):
    st.markdown(T["rules_md"])

split_keyword = (st.text_input(T["split_keyword_label"], value="hangtag").strip().lower() or "hangtag")

st.subheader(T["step1_header"])
bs_files = st.file_uploader(T["bs_uploader_label"], type=["xlsx"], accept_multiple_files=True, key="bs_files")

st.subheader(T["step2_header"])
fr_file = st.file_uploader(T["fr_uploader_label"], type=["xlsx"], key="fr_file")

fr_sheet_name = None
if fr_file is not None:
    try:
        fr_file.seek(0)
        probe_wb = openpyxl.load_workbook(io.BytesIO(fr_file.read()), read_only=True)
        # hide this tool's own internal bookkeeping sheet (tracks which
        # styles it has already claimed) from the sheet picker
        sheet_names = [s for s in probe_wb.sheetnames if s != engine.META_SHEET_NAME]
        fr_file.seek(0)
        if len(sheet_names) == 1:
            fr_sheet_name = sheet_names[0]
        else:
            fr_sheet_name = st.selectbox(T["sheet_select_label"], sheet_names)
    except Exception as e:
        st.error(T["fr_read_error"].format(e=e))

process_clicked = st.button(
    T["process_button"], type="primary", disabled=not (bs_files and fr_file and fr_sheet_name)
)

if process_clicked:
    with st.spinner(T["spinner"]):
        try:
            fr_file.seek(0)
            fr_wb = openpyxl.load_workbook(io.BytesIO(fr_file.read()))

            bs_wbs = []
            for f in bs_files:
                f.seek(0)
                bs_wbs.append(openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True))

            results = engine.process(fr_wb, fr_sheet_name, bs_wbs, split_keyword=split_keyword)

            buf = io.BytesIO()
            fr_wb.save(buf)
            buf.seek(0)

            st.session_state["result_bytes"] = buf.getvalue()
            st.session_state["results"] = results
            st.session_state["fr_filename"] = fr_file.name
        except engine.EngineError as e:
            st.error(T["process_error"].format(e=status_text(str(e))))
            st.session_state.pop("result_bytes", None)
        except Exception as e:
            st.error(T["process_error"].format(e=e))
            st.session_state.pop("result_bytes", None)

if "results" in st.session_state:
    st.subheader(T["results_header"])
    for r in st.session_state["results"]:
        status = r.get("status", "")
        style = r.get("style", r.get("sheet"))
        if status == "ok":
            added = T["rows_added_suffix"].format(n=r["rows_added"]) if r.get("rows_added") else ""
            preserved = (
                T["rows_preserved_suffix"].format(n=r["rows_preserved"])
                if r.get("preserved_existing") and r.get("rows_preserved")
                else ""
            )
            st.success(
                T["result_ok"].format(
                    style=style, lines=r["lines"], start=r["block_start"], end=r["block_end"],
                    added=added, preserved=preserved,
                )
            )
        else:
            st.warning(T["result_warn"].format(style=style, sheet=r.get("sheet"), status=status_text(status)))

if "result_bytes" in st.session_state:
    base_name = st.session_state.get("fr_filename", "FR_updated.xlsx")
    if base_name.lower().endswith(".xlsx"):
        base_name = base_name[:-5]
    out_name = f"{base_name}_TRIM_UPDATED_{datetime.date.today().isoformat()}.xlsx"
    st.download_button(
        T["download_button"],
        data=st.session_state["result_bytes"],
        file_name=out_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )
